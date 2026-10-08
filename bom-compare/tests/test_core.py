"""Test della logica di business (eseguibili con: python -m pytest tests)."""
from __future__ import annotations

import math
import time
from pathlib import Path

import pandas as pd
import pytest

from bomcompare.core.comparator import CompareSettings, compare_all, compare_pair, summarize
from bomcompare.core.config import AppConfig
from bomcompare.core.history import HistoryStore
from bomcompare.core.loader import (BOMLoadError, RawTable, build_bom, detect_columns,
                                    detect_header_row, inspect_file, parse_level, parse_number)
from bomcompare.core.models import ADDED, MODIFIED, REMOVED, UNCHANGED, BOMData
from bomcompare.core.project import ProjectState, load_project, save_project
from bomcompare.export.email_text import build_email_text
from bomcompare.export.excel_report import export_excel

ROOT = Path(__file__).resolve().parents[1]
SAMPLES = ROOT / "sample_data"


@pytest.fixture(scope="module")
def cfg(tmp_path_factory) -> AppConfig:
    return AppConfig(ROOT / "config.json", tmp_path_factory.mktemp("user"))


def make_bom(rows: list[dict], role: str = "X", mapping: dict | None = None) -> BOMData:
    """Crea una BOMData da righe con chiavi = nomi colonna del 'file'."""
    table = pd.DataFrame(rows)
    mapping = mapping or {c: c for c in table.columns}
    rt = RawTable(path="", raw=pd.concat([pd.DataFrame([list(table.columns)]),
                                          pd.DataFrame(table.values)], ignore_index=True))
    rt.mapping = mapping
    return build_bom(rt, role)


# --------------------------------------------------------------------------- #
# Parsing
# --------------------------------------------------------------------------- #
@pytest.mark.parametrize("value,expected", [
    ("1.234,5", 1234.5), ("1,5", 1.5), ("1,234.5", 1234.5), ("€ 12", 12.0),
    (3, 3.0), ("10 pz", 10.0), ("1.000.000", 1_000_000.0), ("-2,5", -2.5),
])
def test_parse_number(value, expected):
    assert parse_number(value) == pytest.approx(expected)


def test_parse_number_invalid():
    assert math.isnan(parse_number("abc"))
    assert math.isnan(parse_number(None))


@pytest.mark.parametrize("value,expected", [("2", 2), (3.0, 3), ("..2", 2), ("...", 3)])
def test_parse_level(value, expected):
    assert parse_level(value) == expected


def test_detect_columns_italian_headers(cfg):
    cols = ["Pos.", "Cod. Art.", "Descrizione articolo", "Q.tà", "U.M.", "Rev",
            "Fornitore", "Costo unitario", "Codice padre"]
    m = detect_columns(cols, cfg.column_aliases())
    assert m["code"] == "Cod. Art."
    assert m["quantity"] == "Q.tà"
    assert m["description"] == "Descrizione articolo"
    assert m["uom"] == "U.M."
    assert m["revision"] == "Rev"
    assert m["price"] == "Costo unitario"
    assert m["parent"] == "Codice padre"


def test_header_row_detection_skips_titles(cfg):
    raw = pd.DataFrame([["DISTINTA MACCHINA X", None, None],
                        [None, "stampata il 01/01", None],
                        ["Codice", "Descrizione", "Qtà"],
                        ["A1", "Vite", 3]])
    assert detect_header_row(raw, cfg.column_aliases()) == 2


# --------------------------------------------------------------------------- #
# Caricamento file di esempio
# --------------------------------------------------------------------------- #
def test_load_sample_xlsx_and_csv(cfg):
    aliases = cfg.column_aliases()
    xlsx = build_bom(inspect_file(str(SAMPLES / "distinta_2_produzione.xlsx"), aliases), "P")
    csv = build_bom(inspect_file(str(SAMPLES / "distinta_2_produzione.csv"), aliases), "P")
    assert xlsx.row_count == csv.row_count > 50
    merged = xlsx.df.merge(csv.df, on="key")
    assert (merged["quantity_x"] == merged["quantity_y"]).all()
    assert (merged["price_x"].fillna(-1) == merged["price_y"].fillna(-1)).all()


def test_multi_sheet_workbook(cfg):
    rt = inspect_file(str(SAMPLES / "distinte_esempio.xlsx"), cfg.column_aliases(),
                      sheet="Progettazione")
    assert rt.sheets == ["Vecchia", "Produzione", "Progettazione"]
    assert rt.is_mapping_complete


def test_load_errors(tmp_path, cfg):
    with pytest.raises(BOMLoadError, match="non esiste"):
        inspect_file(str(tmp_path / "manca.xlsx"), cfg.column_aliases())
    bad = tmp_path / "file.pdf"
    bad.write_text("x")
    with pytest.raises(BOMLoadError, match="non supportato"):
        inspect_file(str(bad), cfg.column_aliases())
    fake = tmp_path / "rotto.xlsx"
    fake.write_text("non sono un excel")
    with pytest.raises(BOMLoadError):
        inspect_file(str(fake), cfg.column_aliases())


def test_codes_keep_leading_zeros_and_numeric_codes():
    bom = make_bom([{"code": "001", "quantity": 1}, {"code": 1001.0, "quantity": 2}])
    assert list(bom.df["code"]) == ["001", "1001"]


def test_duplicates_are_aggregated_and_parents_derived():
    rows = [
        {"level": 1, "code": "ASS1", "quantity": 1},
        {"level": 2, "code": "V1", "quantity": 4},
        {"level": 1, "code": "ASS2", "quantity": 1},
        {"level": 2, "code": "V1", "quantity": 6},
        {"level": 2, "code": "", "quantity": 9},
    ]
    bom = make_bom(rows)
    v1 = bom.df.set_index("key").loc["V1"]
    assert v1["quantity"] == 10
    assert v1["occurrences"] == 2
    assert v1["parent"] == "ASS1 | ASS2"
    assert any("senza codice" in w for w in bom.warnings)


# --------------------------------------------------------------------------- #
# Confronto
# --------------------------------------------------------------------------- #
OLD = [
    {"code": "A", "description": "Vite", "quantity": 10, "revision": "A", "supplier": "S1", "price": 1.0},
    {"code": "B", "description": "Dado", "quantity": 5, "revision": "A", "supplier": "S1", "price": 0.5},
    {"code": "C", "description": "Rondella", "quantity": 2, "revision": "A", "supplier": "S1", "price": 0.1},
    {"code": "D", "description": "Motore", "quantity": 1, "revision": "A", "supplier": "S2", "price": 100.0},
]
NEW = [
    {"code": "A", "description": "VITE", "quantity": 10, "revision": "A", "supplier": "S1", "price": 1.0},
    {"code": "B", "description": "Dado", "quantity": 8, "revision": "A", "supplier": "S1", "price": 0.5},
    {"code": "D", "description": "Motore", "quantity": 1, "revision": "B", "supplier": "S3", "price": 100.0},
    {"code": "E", "description": "Sensore", "quantity": 2, "revision": "A", "supplier": "S2", "price": 20.0},
]


def test_compare_pair_statuses():
    res = compare_pair(make_bom(OLD), make_bom(NEW), CompareSettings(tolerance_pct=10))
    st = dict(zip(res["key"], res["status"]))
    assert st == {"A": UNCHANGED, "B": MODIFIED, "C": REMOVED, "D": MODIFIED, "E": ADDED}
    b = res.set_index("key").loc["B"]
    assert b["delta_qty"] == 3 and b["delta_pct"] == pytest.approx(60)
    assert bool(b["out_of_tol"])
    d = res.set_index("key").loc["D"]
    assert bool(d["chg_revision"]) and bool(d["chg_supplier"])
    assert "Rev. A → B" in d["changes"]
    # ordine "diff": C (eliminato) resta dopo B
    assert list(res["key"]) == ["A", "B", "C", "D", "E"]
    s = summarize(res)
    assert (s["added"], s["removed"], s["modified"], s["unchanged"]) == (1, 1, 2, 1)
    # impatto: B +3*0.5, C -2*0.1, E +2*20
    assert s["impact"] == pytest.approx(1.5 - 0.2 + 40)


def test_compare_fields_setting_limits_modified():
    res = compare_pair(make_bom(OLD), make_bom(NEW),
                       CompareSettings(compare_fields=["quantity"]))
    st = dict(zip(res["key"], res["status"]))
    assert st["D"] == UNCHANGED and st["B"] == MODIFIED


def test_description_case_sensitivity():
    res = compare_pair(make_bom(OLD), make_bom(NEW),
                       CompareSettings(ignore_description_case=False))
    assert dict(zip(res["key"], res["status"]))["A"] == MODIFIED


def test_structure_change_detected():
    old = [{"level": 1, "code": "G1", "quantity": 1}, {"level": 2, "code": "X", "quantity": 1},
           {"level": 1, "code": "G2", "quantity": 1}]
    new = [{"level": 1, "code": "G1", "quantity": 1}, {"level": 1, "code": "G2", "quantity": 1},
           {"level": 2, "code": "X", "quantity": 1}]
    res = compare_pair(make_bom(old), make_bom(new), CompareSettings())
    x = res.set_index("key").loc["X"]
    assert x["status"] == MODIFIED and bool(x["chg_structure"])
    assert "padre G1 → G2" in x["changes"]


def test_compare_all_three_boms(cfg):
    aliases = cfg.column_aliases()
    files = ["distinta_1_vecchia.xlsx", "distinta_2_produzione.xlsx",
             "distinta_3_progettazione.xlsx"]
    boms = [build_bom(inspect_file(str(SAMPLES / f), aliases), f) for f in files]
    result = compare_all(boms, CompareSettings())
    assert set(result.pairs) == {(0, 1), (1, 2), (0, 2)}
    assert result.pair_keys()[:2] == [(0, 1), (1, 2)]
    al = result.aligned
    assert len(al) == al["key"].nunique()
    row = al.set_index("key").loc["FOT-5540"]
    assert row["status_2"] == MODIFIED and bool(row["chg_structure_2"])
    assert al["any_added"].sum() > 0 and al["any_removed"].sum() > 0


def test_performance_large_bom():
    n = 5000
    old = [{"code": f"C{i:05d}", "description": f"Articolo {i}", "quantity": i % 7 + 1,
            "price": 1.5} for i in range(n)]
    new = [dict(r) for r in old[100:]] + [{"code": f"N{i}", "description": "Nuovo",
                                           "quantity": 1, "price": 2.0} for i in range(200)]
    for r in new[::50]:
        r["quantity"] += 1
    a, b = make_bom(old), make_bom(new)
    t0 = time.perf_counter()
    result = compare_all([a, b, b], CompareSettings())
    elapsed = time.perf_counter() - t0
    s = result.summary((0, 1))
    assert s["removed"] == 100 and s["added"] == 200
    assert elapsed < 5, f"confronto troppo lento: {elapsed:.2f}s"


# --------------------------------------------------------------------------- #
# Persistenza ed export
# --------------------------------------------------------------------------- #
def test_project_roundtrip(tmp_path):
    bom = make_bom([{"code": "007", "quantity": 1.5, "description": "X"}], role="Vecchia")
    state = ProjectState(machine="M1", roles=["Vecchia", "Nuova", "Terza"],
                         boms=[bom, None, None], notes={"007": "controllare"})
    state.log_event("Test", "dettaglio")
    path = str(tmp_path / "p.bomcompare")
    save_project(path, state)
    loaded = load_project(path)
    assert loaded.machine == "M1"
    assert loaded.boms[1] is None
    assert loaded.boms[0].df["code"].tolist() == ["007"]
    assert loaded.boms[0].df["quantity"].tolist() == [1.5]
    assert loaded.notes == {"007": "controllare"}
    assert loaded.change_log[0]["event"] == "Test"


def test_history_store(tmp_path):
    store = HistoryStore(tmp_path / "h.sqlite")
    bom = make_bom(OLD, role="Vecchia")
    assert store.add_snapshot(bom, "M1") is not None
    assert store.add_snapshot(bom, "M1") is None  # duplicato identico ignorato
    snaps = store.list_snapshots()
    assert len(snaps) == 1
    again = store.get_snapshot(snaps[0]["id"])
    assert again.df["key"].tolist() == bom.df["key"].tolist()
    cid = store.add_comparison("t", "M1", ["a", "b"], {"x": {"added": 1}}, {"format": "x"})
    assert store.get_comparison_project(cid) == {"format": "x"}


def test_exports(tmp_path):
    result = compare_all([make_bom(OLD, role="Vecchia"), make_bom(NEW, role="Nuova")],
                         CompareSettings())
    xlsx = tmp_path / "r.xlsx"
    export_excel(str(xlsx), result, {"B": "nota"}, [], "M1")
    from openpyxl import load_workbook
    wb = load_workbook(xlsx)
    assert wb.sheetnames[:2] == ["Riepilogo", "Vista allineata"]
    assert "Log modifiche" in wb.sheetnames
    text = build_email_text(result, (0, 1), "M1", {"B": "nota"})
    for section in ("COMPONENTI AGGIUNTI", "COMPONENTI ELIMINATI", "COMPONENTI MODIFICATI",
                    "Nota: nota", "Fornitore S2 → S3"):
        assert section in text
