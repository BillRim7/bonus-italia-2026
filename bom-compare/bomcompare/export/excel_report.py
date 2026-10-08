"""Report Excel con colori mantenuti (openpyxl)."""
from __future__ import annotations

import logging
import math
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from ..core.comparator import ComparisonResult
from ..core.models import ADDED, MODIFIED, REMOVED, UNCHANGED, now_str

log = logging.getLogger(__name__)

# Colori del report (sempre tema chiaro, adatti alla stampa)
FILLS = {
    ADDED: PatternFill("solid", start_color="DCFCE7"),
    REMOVED: PatternFill("solid", start_color="FEE2E2"),
    MODIFIED: PatternFill("solid", start_color="DBEAFE"),
    UNCHANGED: PatternFill("solid", start_color="F3F4F6"),
    "": PatternFill("solid", start_color="FAFAFA"),
}
FONTS = {
    ADDED: Font(color="166534"),
    REMOVED: Font(color="991B1B"),
    MODIFIED: Font(color="1E40AF"),
    UNCHANGED: Font(color="6B7280"),
    "": Font(color="BBBBBB"),
}
CHANGED_FONT = Font(color="1E3A8A", bold=True)
CHANGED_FILL = PatternFill("solid", start_color="BFDBFE")
WARN_FILL = PatternFill("solid", start_color="FDE68A")
HEADER_FILL = PatternFill("solid", start_color="1F2937")
HEADER_FONT = Font(color="FFFFFF", bold=True)
GROUP_FILLS = [PatternFill("solid", start_color=c) for c in
               ("6B7280", "2563EB", "7C3AED", "0D9488", "C2410C", "BE185D")]
TITLE_FONT = Font(size=16, bold=True, color="111827")
THIN = Side(style="thin", color="E5E7EB")
BORDER = Border(bottom=THIN)
NUM_FMT = "#,##0.###"
PCT_FMT = '+0.0"%";-0.0"%";0"%"'
MONEY_FMT = '#,##0.00 "€"'
DELTA_FMT = "+#,##0.###;-#,##0.###;0"


def _v(value: Any) -> Any:
    """Valore adatto a openpyxl (NaN -> None, numpy -> python)."""
    if value is None:
        return None
    if hasattr(value, "item"):
        value = value.item()
    if isinstance(value, float) and math.isnan(value):
        return None
    if isinstance(value, str) and value == "":
        return None
    return value


def _header(ws, row: int, headers: list[str], fill=HEADER_FILL) -> None:
    for col, text in enumerate(headers, start=1):
        c = ws.cell(row=row, column=col, value=text)
        c.fill = fill
        c.font = HEADER_FONT
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 30


def _widths(ws, widths: list[int]) -> None:
    for col, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = w


def export_excel(path: str, result: ComparisonResult, notes: dict[str, str],
                 change_log: list[dict[str, str]], machine: str = "") -> None:
    """Crea il report Excel completo."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Riepilogo"
    _summary_sheet(ws, result, machine)
    _aligned_sheet(wb.create_sheet("Vista allineata"), result, notes)
    used: set[str] = set()
    for pair in result.pair_keys():
        i, j = pair
        name = f"{i + 1}vs{j + 1} {result.names[i][:10]}-{result.names[j][:10]}"
        name = "".join(ch for ch in name if ch not in "[]:*?/\\")[:31]
        while name in used:
            name = name[:29] + f"_{len(used)}"
        used.add(name)
        _pair_sheet(wb.create_sheet(name), result, pair, notes)
    _log_sheet(wb.create_sheet("Log modifiche"), change_log)
    wb.save(path)
    log.info("Report Excel salvato: %s", path)


# --------------------------------------------------------------------------- #
def _summary_sheet(ws, result: ComparisonResult, machine: str) -> None:
    ws["A1"] = "Report confronto distinte base"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = f"Macchina: {machine or '-'}    Generato il: {now_str()}"
    ws["A2"].font = Font(color="6B7280")

    ws["A4"] = "Distinte confrontate"
    ws["A4"].font = Font(bold=True, size=12)
    _header(ws, 5, ["#", "Ruolo", "File", "Foglio", "Codici", "Caricata il"])
    for i, b in enumerate(result.boms):
        row = 6 + i
        for col, v in enumerate([i + 1, b.role, b.file_name, b.sheet or "",
                                 b.row_count, b.loaded_at], start=1):
            ws.cell(row=row, column=col, value=v).border = BORDER
        ws.cell(row=row, column=2).fill = GROUP_FILLS[i % len(GROUP_FILLS)]
        ws.cell(row=row, column=2).font = Font(color="FFFFFF", bold=True)

    start = 7 + len(result.boms)
    ws.cell(row=start, column=1, value="Statistiche per confronto").font = Font(bold=True, size=12)
    headers = ["Confronto", "Aggiunti", "Eliminati", "Modificati", "Invariati",
               "Cambi quantità", "Cambi revisione", "Cambi fornitore", "Cambi prezzo",
               "Spostamenti struttura", "Fuori tolleranza", "Costo iniziale",
               "Costo finale", "Impatto economico"]
    _header(ws, start + 1, headers)
    for r, pair in enumerate(result.pair_keys(), start=start + 2):
        s = result.summary(pair)
        values = [result.pair_label(pair), s["added"], s["removed"], s["modified"],
                  s["unchanged"], s["qty_changes"], s["revision_changes"],
                  s["supplier_changes"], s["price_changes"], s["structure_changes"],
                  s["out_of_tolerance"], s["cost_a"], s["cost_b"], s["impact"]]
        for col, v in enumerate(values, start=1):
            c = ws.cell(row=r, column=col, value=_v(v))
            c.border = BORDER
            if col >= 12:
                c.number_format = MONEY_FMT
        for col, st in ((2, ADDED), (3, REMOVED), (4, MODIFIED), (5, UNCHANGED)):
            ws.cell(row=r, column=col).fill = FILLS[st]
            ws.cell(row=r, column=col).font = FONTS[st]

    leg = start + 3 + len(result.pairs)
    ws.cell(row=leg, column=1, value="Legenda").font = Font(bold=True, size=12)
    legend = [(ADDED, "Componente aggiunto nella versione più recente"),
              (REMOVED, "Componente eliminato nella versione più recente"),
              (MODIFIED, "Stesso codice, dati diversi (quantità, revisione, fornitore…)"),
              (UNCHANGED, "Nessuna differenza")]
    for k, (st, txt) in enumerate(legend, start=leg + 1):
        c = ws.cell(row=k, column=1, value=st)
        c.fill, c.font = FILLS[st], FONTS[st]
        ws.cell(row=k, column=2, value=txt)
    k = leg + len(legend) + 1
    c = ws.cell(row=k, column=1, value="Fuori toll.")
    c.fill = WARN_FILL
    ws.cell(row=k, column=2,
            value=f"Variazione quantità oltre ±{result.settings.tolerance_pct:g}%")
    _widths(ws, [34, 26, 34, 14, 12, 20, 14, 14, 14, 14, 16, 16, 16, 18])


# --------------------------------------------------------------------------- #
def _aligned_sheet(ws, result: ComparisonResult, notes: dict[str, str]) -> None:
    al = result.aligned
    n = result.n
    has_price = result.any_field("price")
    has_level = result.any_field("level")
    per_bom = ["Qtà", "U.M.", "Rev.", "Fornitore"] + (["Prezzo"] if has_price else [])
    per_step = ["Stato", "Δ Qtà", "Δ %"]

    # Riga 1: gruppi colorati con il nome della distinta
    headers = ["Codice", "Descrizione"] + (["Liv."] if has_level else [])
    fixed = len(headers)
    group_spans: list[tuple[int, int, str, int]] = []
    for i in range(n):
        start = len(headers) + 1
        headers += per_bom
        if i > 0:
            headers += per_step
        group_spans.append((start, len(headers), result.names[i], i))
    headers.append("Note")
    for start, end, name, i in group_spans:
        ws.merge_cells(start_row=1, start_column=start, end_row=1, end_column=end)
        c = ws.cell(row=1, column=start, value=name)
        c.fill = GROUP_FILLS[i % len(GROUP_FILLS)]
        c.font = HEADER_FONT
        c.alignment = Alignment(horizontal="center")
    _header(ws, 2, headers)

    cols = {c: al[c].tolist() for c in al.columns}
    for r in range(len(al)):
        row = r + 3
        status_row = [cols[f"status_{k}"][r] if k > 0 else "" for k in range(n)]
        values: list[Any] = [cols["code"][r], cols["description"][r]]
        if has_level:
            lvl = next((cols[f"level_{i}"][r] for i in range(n - 1, -1, -1)
                        if cols[f"present_{i}"][r]), None)
            values.append(_v(lvl))
        fmt_cells: list[tuple[int, str, bool, str | None]] = []  # col, status, changed, numfmt
        for i in range(n):
            present = cols[f"present_{i}"][r]
            # Colore: per la prima distinta lo stato verso la successiva.
            st = status_row[i] if i > 0 else (cols["status_1"][r] if n > 1 else "")
            if not present:
                st_panel = ""
            else:
                st_panel = st
            k = i if i > 0 else 1
            chg = lambda f: bool(cols.get(f"chg_{f}_{k}", [False] * len(al))[r]) and present
            base = len(values) + 1
            vals = [cols[f"quantity_{i}"][r], cols[f"uom_{i}"][r], cols[f"revision_{i}"][r],
                    cols[f"supplier_{i}"][r]]
            flags = [chg("quantity"), chg("uom"), chg("revision"), chg("supplier")]
            fmts = [NUM_FMT, None, None, None]
            if has_price:
                vals.append(cols[f"price_{i}"][r])
                flags.append(chg("price"))
                fmts.append(MONEY_FMT)
            if i > 0:
                vals += [status_row[i], cols[f"delta_qty_{i}"][r], cols[f"delta_pct_{i}"][r]]
                flags += [False, False, bool(cols[f"out_of_tol_{i}"][r])]
                fmts += [None, DELTA_FMT, PCT_FMT]
            for off, (v, fl, nf) in enumerate(zip(vals, flags, fmts)):
                values.append(_v(v) if present or off >= len(per_bom) else None)
                fmt_cells.append((base + off, st_panel if (present or off >= len(per_bom)) else "",
                                  fl, nf))
        values.append(notes.get(cols["key"][r], ""))
        for col, v in enumerate(values, start=1):
            ws.cell(row=row, column=col, value=_v(v))
        for col, st, changed, nf in fmt_cells:
            c = ws.cell(row=row, column=col)
            c.fill = CHANGED_FILL if changed and st == MODIFIED else FILLS.get(st, FILLS[""])
            c.font = CHANGED_FONT if changed and st == MODIFIED else FONTS.get(st, FONTS[""])
            if nf:
                c.number_format = nf
            if nf == PCT_FMT and changed:
                c.fill = WARN_FILL
        # Codice/descrizione colorati con lo stato più "forte" della riga
        overall = (REMOVED if cols["any_removed"][r] else ADDED if cols["any_added"][r]
                   else MODIFIED if cols["any_modified"][r] else UNCHANGED)
        for col in range(1, fixed + 1):
            ws.cell(row=row, column=col).fill = FILLS[overall]
    # Δ% è memorizzato come numero (es. 12.5): il formato aggiunge il simbolo
    ws.freeze_panes = ws.cell(row=3, column=fixed + 1)
    ws.auto_filter.ref = f"A2:{get_column_letter(len(headers))}{len(al) + 2}"
    widths = [16, 40] + ([6] if has_level else [])
    for i in range(n):
        widths += [9, 6, 6, 22] + ([11] if has_price else [])
        if i > 0:
            widths += [12, 9, 9]
    widths.append(40)
    _widths(ws, widths)


# --------------------------------------------------------------------------- #
def _pair_sheet(ws, result: ComparisonResult, pair: tuple[int, int],
                notes: dict[str, str]) -> None:
    i, j = pair
    df = result.pairs[pair]
    a, b = result.names[i], result.names[j]
    has_price = result.boms[i].has("price") or result.boms[j].has("price")
    has_level = result.boms[i].has("level") or result.boms[j].has("level")
    ws["A1"] = f"Confronto: {a}  →  {b}"
    ws["A1"].font = Font(size=13, bold=True)
    spec: list[tuple[str, str, str | None, str | None]] = [
        ("Stato", "status", None, None),
        ("Codice", "code", None, None),
        ("Descrizione", "description", None, "description"),
    ]
    if has_level:
        spec += [(f"Liv. {a[:12]}", "level_a", None, "structure"),
                 (f"Liv. {b[:12]}", "level_b", None, "structure"),
                 (f"Padre {b[:12]}", "parent_b", None, "structure")]
    spec += [
        (f"Qtà {a[:12]}", "quantity_a", NUM_FMT, None),
        (f"Qtà {b[:12]}", "quantity_b", NUM_FMT, "quantity"),
        ("Δ Qtà", "delta_qty", DELTA_FMT, "quantity"),
        ("Δ %", "delta_pct", PCT_FMT, None),
        (f"U.M. {a[:12]}", "uom_a", None, None),
        (f"U.M. {b[:12]}", "uom_b", None, "uom"),
        (f"Rev. {a[:12]}", "revision_a", None, None),
        (f"Rev. {b[:12]}", "revision_b", None, "revision"),
        (f"Fornitore {a[:12]}", "supplier_a", None, None),
        (f"Fornitore {b[:12]}", "supplier_b", None, "supplier"),
    ]
    if has_price:
        spec += [(f"Prezzo {a[:12]}", "price_a", MONEY_FMT, None),
                 (f"Prezzo {b[:12]}", "price_b", MONEY_FMT, "price"),
                 ("Impatto €", "impact", MONEY_FMT, None)]
    spec += [("Fuori toll.", "out_of_tol", None, None),
             ("Modifiche", "changes", None, None), ("Note", "_note", None, None)]
    _header(ws, 3, [s[0] for s in spec])
    cols = {c: df[c].tolist() for c in df.columns}
    for r in range(len(df)):
        row = r + 4
        st = cols["status"][r]
        for col, (_h, src, nf, chg_field) in enumerate(spec, start=1):
            if src == "_note":
                v = notes.get(cols["key"][r], "")
            elif src == "out_of_tol":
                v = "SÌ" if cols["out_of_tol"][r] else ""
            else:
                v = cols[src][r]
            c = ws.cell(row=row, column=col, value=_v(v))
            changed = bool(chg_field and cols.get(f"chg_{chg_field}", [False] * len(df))[r])
            c.fill = CHANGED_FILL if changed else FILLS[st]
            c.font = CHANGED_FONT if changed else FONTS[st]
            if nf:
                c.number_format = nf
            if src in ("delta_pct", "out_of_tol") and cols["out_of_tol"][r]:
                c.fill = WARN_FILL
    ws.freeze_panes = "D4"
    ws.auto_filter.ref = f"A3:{get_column_letter(len(spec))}{len(df) + 3}"
    widths = {"Stato": 12, "Codice": 16, "Descrizione": 40, "Modifiche": 42, "Note": 36}
    _widths(ws, [widths.get(s[0], 12 if not s[0].startswith("Fornitore") else 22)
                 for s in spec])


def _log_sheet(ws, change_log: list[dict[str, str]]) -> None:
    _header(ws, 1, ["Data e ora", "Evento", "Dettaglio"])
    for r, e in enumerate(change_log, start=2):
        ws.cell(row=r, column=1, value=e.get("timestamp", ""))
        ws.cell(row=r, column=2, value=e.get("event", ""))
        ws.cell(row=r, column=3, value=e.get("detail", ""))
    _widths(ws, [20, 30, 90])
